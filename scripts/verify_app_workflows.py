#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""End-to-end API workflow verification for FinDataHub."""
import argparse
import json
import time
import uuid
from io import BytesIO
from urllib import error, parse, request

from openpyxl import Workbook, load_workbook


FIXED_MODULES = {
    'profit_rate': {
        'payload': {'year': 2091, 'gross_profit_margin': 0, 'net_profit_margin': 0},
        'update': {'gross_profit_margin': 1.25},
        'zero_fields': ['gross_profit_margin', 'net_profit_margin'],
    },
    'non_recurring': {
        'payload': {'year': 2091, 'non_recurring_profit': 0, 'non_recurring_growth': 0},
        'update': {'non_recurring_profit': 1.25},
        'zero_fields': ['non_recurring_profit', 'non_recurring_growth'],
    },
    'roe_net_asset': {
        'payload': {'year': 2091, 'roe': 0, 'net_asset_per_share': 0},
        'update': {'roe': 1.25},
        'zero_fields': ['roe', 'net_asset_per_share'],
    },
    'pe_valuation': {
        'payload': {'year': 2091, 'pe_high': 0, 'pe_mid': 0, 'pe_low': 0, 'eps': 0, 'type': 'actual'},
        'update': {'pe_high': 1.25},
        'zero_fields': ['pe_high', 'pe_mid', 'pe_low', 'eps'],
    },
    'shareholder_structure': {
        'payload': {'stat_date': '2091-06-30', 'shareholder_type': '', 'holding_ratio': 0, 'change_ratio': 0},
        'update': {'holding_ratio': 1.25},
        'zero_fields': ['holding_ratio', 'change_ratio'],
    },
    'shareholder_count': {
        'payload': {'stat_date': '2091-06-30', 'total_holders': 0, 'change': 0},
        'update': {'change': 1.25},
        'zero_fields': ['change'],
    },
    'rd_expense': {
        'payload': {'year': 2091, 'revenue': 0, 'rd_expense': 0, 'rd_ratio': 0, 'rd_growth': 0, 'rd_return': 0},
        'update': {'revenue': 1.25},
        'zero_fields': ['revenue', 'rd_expense', 'rd_ratio', 'rd_growth', 'rd_return'],
    },
    'rd_staff': {
        'payload': {
            'year': 2091, 'staff_count': 0, 'growth': 0, 'percent_of_total': 0,
            'bachelor': 0, 'master': 0, 'bachelor_master_ratio': 0, 'remark': ''
        },
        'update': {'growth': 1.25},
        'zero_fields': ['growth', 'percent_of_total', 'bachelor_master_ratio'],
    },
}


class Client:
    def __init__(self, base_url):
        self.base_url = base_url.rstrip('/')

    def request(self, method, path, data=None, expected=200, raw=False, headers=None):
        body = None
        request_headers = headers.copy() if headers else {}
        if data is not None and not isinstance(data, (bytes, bytearray)):
            body = json.dumps(data).encode('utf-8')
            request_headers['Content-Type'] = 'application/json'
        elif data is not None:
            body = data

        req = request.Request(
            self.base_url + path,
            data=body,
            headers=request_headers,
            method=method,
        )
        try:
            with request.urlopen(req, timeout=20) as resp:
                payload = resp.read()
                status = resp.status
                response_headers = resp.headers
        except error.HTTPError as exc:
            payload = exc.read()
            status = exc.code
            response_headers = exc.headers

        if status != expected:
            text = payload.decode('utf-8', errors='replace')
            raise AssertionError(f'{method} {path} expected {expected}, got {status}: {text}')
        if raw:
            return payload, response_headers
        if not payload:
            return None
        parsed = json.loads(payload.decode('utf-8'))
        if status == 200 and parsed.get('code') != 200:
            raise AssertionError(f'{method} {path} business failure: {parsed}')
        return parsed.get('data')

    def multipart(self, path, fields, files, expected=200):
        boundary = '----FinDataHubBoundary' + uuid.uuid4().hex
        chunks = []
        for name, value in fields.items():
            chunks.extend([
                f'--{boundary}\r\n'.encode(),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode(),
                str(value).encode('utf-8'),
                b'\r\n',
            ])
        for name, filename, content, content_type in files:
            chunks.extend([
                f'--{boundary}\r\n'.encode(),
                f'Content-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'.encode(),
                f'Content-Type: {content_type}\r\n\r\n'.encode(),
                content,
                b'\r\n',
            ])
        chunks.append(f'--{boundary}--\r\n'.encode())
        return self.request(
            'POST',
            path,
            data=b''.join(chunks),
            expected=expected,
            headers={'Content-Type': f'multipart/form-data; boundary={boundary}'},
        )


def assert_xlsx(payload):
    if not payload.startswith(b'PK'):
        raise AssertionError('expected an XLSX/ZIP payload')
    load_workbook(BytesIO(payload), read_only=True).close()


def workbook_bytes(headers, rows, sheet_name='Sheet1'):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def cleanup(client, cleanup_actions):
    for method, path in reversed(cleanup_actions):
        try:
            client.request(method, path, expected=200)
        except Exception:
            pass


def purge_qa_artifacts(client):
    try:
        templates = client.request('GET', '/api/industry-templates/all')
    except Exception:
        templates = []
    for template in templates:
        code = template.get('code') or ''
        if not code.startswith('qa_industry_'):
            continue
        for _ in range(2):
            try:
                client.request('DELETE', f'/api/industry-templates/{template["id"]}', expected=200)
            except Exception:
                break

    try:
        companies = client.request('GET', '/api/companies?keyword=QA&per_page=1000')
    except Exception:
        companies = {'items': []}
    for company in companies.get('items', []):
        code = company.get('code') or ''
        name = company.get('name') or ''
        if code.startswith('QA') or name.startswith('QA测试公司'):
            try:
                client.request('DELETE', f'/api/companies/{company["id"]}', expected=200)
            except Exception:
                pass


def verify_fixed_modules(client, company_id):
    cleanup_actions = []
    try:
        for module_name, spec in FIXED_MODULES.items():
            client.request('GET', f'/api/{module_name}?page=1&per_page=2')
            payload = {'company_id': company_id, **spec['payload']}
            created = client.request('POST', f'/api/{module_name}', payload)
            cleanup_actions.append(('DELETE', f'/api/{module_name}/{created["id"]}'))

            for field in spec['zero_fields']:
                if created.get(field) not in (0, 0.0):
                    raise AssertionError(f'{module_name}.{field} should preserve zero, got {created.get(field)!r}')

            duplicate = client.request('POST', f'/api/{module_name}', payload, expected=400)
            if duplicate and duplicate.get('code') == 200:
                raise AssertionError(f'{module_name} duplicate insert unexpectedly succeeded')

            updated = client.request('PUT', f'/api/{module_name}/{created["id"]}', spec['update'])
            for field, value in spec['update'].items():
                if updated.get(field) != value:
                    raise AssertionError(f'{module_name}.{field} update mismatch: {updated.get(field)!r}')

            export_payload, _ = client.request('GET', f'/api/{module_name}/export', raw=True)
            assert_xlsx(export_payload)

            template_payload, _ = client.request('GET', f'/api/templates/{module_name}', raw=True)
            assert_xlsx(template_payload)
    finally:
        cleanup(client, cleanup_actions)


def verify_fixed_import_errors(client, company_code, company_name):
    bad_workbook = workbook_bytes(
        ['代码', '个股名称', '年份', '销售毛利率(%)', '销售净利率(%)'],
        [
            [company_code, company_name, '', 1, 2],
            [company_code, company_name, 'bad-year', 1, 2],
        ],
    )
    preview = client.multipart(
        '/api/data_import/upload',
        {'module': 'profit_rate'},
        [('file', 'bad-profit-rate.xlsx', bad_workbook, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')],
    )
    details = preview['preview']['profit_rate']
    if details['success_count'] != 0 or details['error_count'] != 2:
        raise AssertionError(f'fixed import row errors not reported correctly: {details}')


def verify_industry_templates(client, company_id, company_code, company_name, suffix):
    code = f'qa_industry_{suffix}'
    cleanup_actions = []
    template = client.request('POST', '/api/industry-templates', {
        'code': code,
        'name': 'QA行业模板',
        'fields': [
            {'keyword': 'channel_count', 'label': '渠道数量', 'data_type': 'number', 'is_required': True},
            {'keyword': 'launch_date', 'label': '上线日期', 'data_type': 'date', 'is_required': False},
            {'keyword': 'region_note', 'label': '区域备注', 'data_type': 'string', 'is_required': False},
        ],
    })
    cleanup_actions.append(('DELETE', f'/api/industry-templates/{template["id"]}'))
    try:
        locked_templates = client.request('GET', '/api/industry-templates')
        if any(item['code'] == code for item in locked_templates):
            raise AssertionError('draft template should not appear in locked template list')

        client.request('POST', f'/api/industry-data/{code}', {
            'company_id': company_id,
            'year': 2091,
            'channel_count': 1,
        }, expected=409)

        confirmed = client.request('POST', f'/api/industry-templates/{template["id"]}/confirm')
        if not confirmed['is_locked']:
            raise AssertionError('template was not locked after confirm')

        client.request('PUT', f'/api/industry-templates/{template["id"]}', {
            'name': 'QA行业模板',
            'code': code,
            'fields': [{'keyword': 'other_field', 'label': '其他字段', 'data_type': 'string'}],
        }, expected=409)

        template_payload, _ = client.request('GET', f'/api/industry-templates/{code}/template', raw=True)
        assert_xlsx(template_payload)

        created = client.request('POST', f'/api/industry-data/{code}', {
            'company_code': company_code,
            'company_name': company_name,
            'year': 2091,
            'channel_count': 0,
            'launch_date': '',
            'region_note': '华东',
        })
        cleanup_actions.append(('DELETE', f'/api/industry-data/{code}/{created["id"]}'))
        if created.get('channel_count') != 0:
            raise AssertionError('industry data should preserve zero values')

        updated = client.request('PUT', f'/api/industry-data/{code}/{created["id"]}', {
            'region_note': '',
        })
        if updated.get('region_note') is not None:
            raise AssertionError('blank optional industry field should clear to null')

        client.request('POST', f'/api/industry-data/{code}/compare', {
            'company_ids': [company_id],
            'years': [2091],
            'metric': 'missing_metric',
        }, expected=400)

        compare = client.request('POST', f'/api/industry-data/{code}/compare', {
            'company_ids': [company_id],
            'years': [2091],
            'metric': 'channel_count',
        })
        if not compare:
            raise AssertionError('industry compare returned no data')

        export_payload, _ = client.request('GET', f'/api/industry-data/{code}/export', raw=True)
        assert_xlsx(export_payload)

        import_payload = workbook_bytes(
            ['代码', '个股名称', '年份', '渠道数量', '上线日期', '区域备注'],
            [
                [company_code, company_name, 2092, 3, '', ''],
                [company_code, company_name, 2093, 'not-a-number', '', '错误行'],
                ['', company_name, 2094, 4, '', '缺代码'],
            ],
        )
        import_result = client.multipart(
            f'/api/industry-data/{code}/import',
            {},
            [('file', 'industry-import.xlsx', import_payload, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')],
        )
        if import_result['success'] != 1 or import_result['error'] != 2:
            raise AssertionError(f'industry import row-level errors mismatch: {import_result}')

        records = client.request('GET', f'/api/industry-data/{code}?page=1&per_page=20')
        for item in records['items']:
            if item['year'] == 2092:
                cleanup_actions.append(('DELETE', f'/api/industry-data/{code}/{item["id"]}'))

        clone = client.request('POST', f'/api/industry-templates/{template["id"]}/clone', {})
        cleanup_actions.append(('DELETE', f'/api/industry-templates/{clone["id"]}'))
        if clone['is_locked'] or clone['version'] <= confirmed['version']:
            raise AssertionError('clone should be a newer unlocked draft')

        client.request('DELETE', f'/api/industry-templates/{template["id"]}')
        all_templates = client.request('GET', '/api/industry-templates/all')
        disabled = next(item for item in all_templates if item['id'] == template['id'])
        if disabled['is_active']:
            raise AssertionError('template with data should be disabled, not active')
    finally:
        cleanup(client, cleanup_actions)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('base_url', nargs='?', default='http://127.0.0.1:5002')
    args = parser.parse_args()

    client = Client(args.base_url)
    suffix = str(int(time.time()))
    company_code = f'QA{suffix[-8:]}'
    company_name = f'QA测试公司{suffix[-4:]}'
    cleanup_actions = []

    try:
        client.request('GET', '/health', raw=True)
        client.request('GET', '/', raw=True)
        purge_qa_artifacts(client)
        company = client.request('POST', '/api/companies', {
            'code': company_code,
            'name': company_name,
            'business': 'workflow verification',
        })
        cleanup_actions.append(('DELETE', f'/api/companies/{company["id"]}'))

        verify_fixed_modules(client, company['id'])
        verify_fixed_import_errors(client, company_code, company_name)
        verify_industry_templates(client, company['id'], company_code, company_name, suffix)

        print(f'FinDataHub workflows verified at {args.base_url}')
    finally:
        cleanup(client, cleanup_actions)
        purge_qa_artifacts(client)


if __name__ == '__main__':
    main()
